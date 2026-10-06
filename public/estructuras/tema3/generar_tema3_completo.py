# -*- coding: utf-8 -*-
"""
QFDOS - Generador Maestro de Estructuras y Propiedades: Tema 3
==============================================================
Genera:
  - estructuras_qfdos.csv
  - propiedades_qfdos.csv
  - estructuras_qfdos.xlsx (con hojas Propiedades y Estructuras con imágenes)
  - Imágenes individuales PNG de cada estructura
  - Figuras compuestas y de correlación estructura-actividad (SAR)
"""

import os
import sys
import csv
import re

from rdkit import Chem, RDLogger
from rdkit.Chem import (AllChem, Crippen, Descriptors, Draw, Lipinski,
                        rdMolDescriptors)
from rdkit.Chem.Draw import rdMolDraw2D
from rdkit.Chem import QED

RDLogger.DisableLog('rdApp.*')

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

TEMA3_DIR = os.path.dirname(os.path.abspath(__file__))
IMGDIR = os.path.join(TEMA3_DIR, 'img')
os.makedirs(IMGDIR, exist_ok=True)


# ─── 1. COMPUESTOS DEL TEMA 3 (Dopamina) ───

COMPUESTOS_TEMA3 = [{'id': 'QFDOS-127', 'nombre': 'Epinina (N-metildopamina)', 'nombre_en': 'Epinine (deoxyepinephrine)', 'tema': 'Analogos restringidos de la dopamina', 'familia': 'Catecolamina', 'smiles': 'CNCCc1ccc(O)c(O)c1', 'cas': '501-15-5', 'notas': 'N-metildopamina (desoxiadrenalina). Conserva el catecol y la amina básica de la dopamina con libre giro en la cadena etilamínica: muchas conformaciones accesibles. Referencia flexible frente a la apomorfina, que contiene su esqueleto congelado. El metilo en N es el mismo que distingue adrenalina de noradrenalina. Es el metabolito activo del profármaco ibopamina.', 'smiles_ph74': 'C[NH2+]CCc1ccc(O)c(O)c1', 'chembl_id': ''},
    {'id': 'QFDOS-128', 'nombre': 'ADTN (2-amino-6,7-dihidroxitetralina)', 'nombre_en': 'ADTN (6,7-dihydroxy-2-aminotetralin)', 'tema': 'Analogos restringidos de la dopamina', 'familia': 'Aminotetralina', 'smiles': 'NC1CCc2cc(O)c(O)cc2C1', 'cas': '', 'notas': '2-Amino-6,7-dihidroxi-1,2,3,4-tetrahidronaftaleno. Agonista dopaminérgico potente: la tetralina fija la cadena etilamínica de la dopamina en conformación antiperiplanar (trans extendida) con el catecol. Se guarda sin estereoquímica (racémico, como en la literatura farmacológica clásica). Junto con la apomorfina es el argumento de que la conformación activa de la dopamina es la antiperiplanar.', 'smiles_ph74': '[NH3+]C1CCc2cc(O)c(O)cc2C1', 'chembl_id': ''},
    {'id': 'QFDOS-129', 'nombre': 'Apomorfina', 'nombre_en': 'Apomorphine', 'tema': 'Analogos restringidos de la dopamina', 'familia': 'Aporfina', 'smiles': 'CN1CCc2cccc3c2[C@H]1Cc1ccc(O)c(O)c1-3', 'cas': '58-00-4', 'notas': 'Agonista D1/D2 obtenido por transposición ácida de la morfina (HCl, deshidratación y migración que aromatiza el anillo C, ver QFDOS-130). Contiene la epinina con la conformación antiperiplanar congelada en el sistema aporfínico. Configuración (6aR). Uso: rescate subcutáneo en los periodos off del Parkinson; emético central potente por estímulo de la zona quimiorreceptora gatillo, coherente con la acción emética de la dopamina que luego bloquea la metoclopramida.', 'smiles_ph74': 'C[NH+]1CCc2cccc3c2[C@H]1Cc1ccc(O)c(O)c1-3', 'chembl_id': ''},
    {'id': 'QFDOS-130', 'nombre': 'Morfina', 'nombre_en': 'Morphine', 'tema': 'Analogos restringidos de la dopamina', 'familia': 'Morfinano', 'smiles': 'CN1CC[C@]23c4c5ccc(O)c4O[C@H]2[C@@H](O)C=C[C@H]3[C@H]1C5', 'cas': '57-27-2', 'notas': 'Material de partida de la apomorfina. En medio ácido fuerte se abre el puente éter 4,5, se pierde agua y el anillo C aromatiza, generando el catecol y el esqueleto aporfínico. Pasa de opioide (receptor mu) a agonista dopaminérgico: el mismo carbono, otro farmacóforo. Cinco estereocentros, configuración natural (−).', 'smiles_ph74': 'C[NH+]1CC[C@]23c4c5ccc(O)c4O[C@H]2[C@@H](O)C=C[C@H]3[C@H]1C5', 'chembl_id': ''},
    {'id': 'QFDOS-131', 'nombre': 'Benserazida', 'nombre_en': 'Benserazide', 'tema': 'Profarmacos de la dopamina', 'familia': 'Hidrazida de serina', 'smiles': 'NC(CO)C(=O)NNCc1ccc(O)c(O)c1O', 'cas': '322-35-0', 'notas': 'Inhibidor de la dopa-descarboxilasa periférica (AADC) asociado a levodopa (Madopar). Hidrazida de la serina con un 2,3,4-trihidroxibencilo: la función hidrazina atrapa el piridoxal fosfato, igual que la carbidopa (QFDOS-094). Muy polar, no cruza la BHE, de modo que solo inhibe la descarboxilación fuera del SNC. Comercializada como racemato (DL-serina); se guarda sin estereoquímica.', 'smiles_ph74': '[NH3+]C(CO)C(=O)NNCc1ccc(O)c(O)c1O', 'chembl_id': ''},
    {'id': 'QFDOS-132', 'nombre': 'Dipivaloildopamina', 'nombre_en': 'Dopamine 3,4-dipivalate', 'tema': 'Profarmacos de la dopamina', 'familia': 'Ester de catecol', 'smiles': 'NCCc1ccc(OC(=O)C(C)(C)C)c(OC(=O)C(C)(C)C)c1', 'cas': '', 'notas': 'Profármaco por latentización del catecol como diéster de ácido piválico. Suficientemente lipófilo para atravesar la BHE, pero no queda retenido en el cerebro y su acción es breve; además, las esterasas plasmáticas liberan dopamina periférica con la consiguiente toxicidad. Contraejemplo que motiva el sistema de inversión de polaridad (QFDOS-133). El terc-butilo del pivalato frena la hidrólisis por impedimento estérico.', 'smiles_ph74': 'CC(C)(C)C(=O)Oc1ccc(CC[NH3+])cc1OC(=O)C(C)(C)C', 'chembl_id': ''},
    {'id': 'QFDOS-133', 'nombre': 'Doble profarmaco dihidropiridinico de dopamina', 'nombre_en': 'Dopamine dihydropyridine chemical delivery system', 'tema': 'Profarmacos de la dopamina', 'familia': '1,4-Dihidropiridina (CDS)', 'smiles': 'CN1C=CCC(C(=O)NCCc2ccc(OC(=O)C(C)(C)C)c(OC(=O)C(C)(C)C)c2)=C1', 'cas': '', 'notas': 'Sistema de liberación química de Bodor. La amina primaria va como amida del ácido 1-metil-1,4-dihidropiridina-3-carboxílico y el catecol como dipivalato: molécula neutra y lipófila que entra en el SNC. Dentro, oxidorreductasas dependientes de NAD(P)+ oxidan la dihidropiridina a piridinio (QFDOS-134), que es catiónico y ya no puede salir por la BHE: inversión de la polaridad. Dos hidrólisis (ésteres y amida) liberan dopamina y trigonelina (QFDOS-135). Lo que se oxida en periferia se elimina rápidamente por ser polar. Resultado: concentración alta en cerebro y baja o nula en plasma.', 'smiles_ph74': 'CN1C=CCC(C(=O)NCCc2ccc(OC(=O)C(C)(C)C)c(OC(=O)C(C)(C)C)c2)=C1', 'chembl_id': ''},
    {'id': 'QFDOS-134', 'nombre': 'Sal de piridinio del profarmaco (forma retenida)', 'nombre_en': 'Dopamine-trigonellinamide pyridinium', 'tema': 'Profarmacos de la dopamina', 'familia': 'Piridinio cuaternario', 'smiles': 'C[n+]1cccc(C(=O)NCCc2ccc(O)c(O)c2)c1', 'cas': '', 'notas': 'Intermedio del doble profármaco tras la oxidación de la dihidropiridina y la hidrólisis de los pivalatos. Catión permanente (amonio aromático cuaternario), por eso queda atrapado tras la BHE: es la forma depósito en el SNC. La hidrólisis de la amida libera dopamina y trigonelina. Se guarda como catión, según el convenio de la base para amonios cuaternarios permanentes.', 'smiles_ph74': 'C[n+]1cccc(C(=O)NCCc2ccc(O)c(O)c2)c1', 'chembl_id': ''},
    {'id': 'QFDOS-135', 'nombre': 'Trigonelina', 'nombre_en': 'Trigonelline', 'tema': 'Profarmacos de la dopamina', 'familia': 'Betaina de piridinio', 'smiles': 'C[n+]1cccc(C(=O)[O-])c1', 'cas': '535-83-1', 'notas': 'N-metilnicotinato (betaína, sal interna). Subproducto del doble profármaco en el SNC: no tóxico, polar, se elimina. Alcaloide natural del café y de la alholva y metabolito de la niacina. Se guarda como zwitterión, mismo criterio que los N-óxidos de la base.', 'smiles_ph74': 'C[n+]1cccc(C(=O)[O-])c1', 'chembl_id': ''},
    {'id': 'QFDOS-136', 'nombre': '(R)-Selegilina', 'nombre_en': '(R)-Selegiline', 'tema': 'Inhibidores MAO-B/COMT y liberadores', 'familia': 'Propargilamina', 'smiles': 'C#CCN(C)[C@H](C)Cc1ccccc1', 'cas': '14611-51-9', 'notas': 'N,α-dimetil-N-propargilfenetilamina. Inhibidor irreversible y selectivo de la MAO-B (inhibidor suicida: el propargilo se oxida y forma un aducto covalente con el FAD). La MAO-B predomina en el estriado y metaboliza allí la mayor parte de la dopamina. Eutómero (R)-(−). Metabolitos: (R)-desmetilselegilina y (R)-metanfetamina, mucho menos estimulante que la (S). Amina terciaria con pKa ≈ 7.5: a pH 7.4 coexisten forma neutra y protonada.', 'smiles_ph74': 'C#CC[NH+](C)[C@H](C)Cc1ccccc1', 'chembl_id': ''},
    {'id': 'QFDOS-137', 'nombre': '(S)-Selegilina', 'nombre_en': '(S)-Selegiline', 'tema': 'Inhibidores MAO-B/COMT y liberadores', 'familia': 'Propargilamina', 'smiles': 'C#CCN(C)[C@@H](C)Cc1ccccc1', 'cas': '', 'notas': 'Distómero de la selegilina. Su N-desalquilación produce (S)-metanfetamina (dextrometanfetamina, QFDOS-068), estimulante central responsable de los efectos indeseados; por eso se comercializa el enantiómero (R). Par enantiomérico con QFDOS-136.', 'smiles_ph74': 'C#CC[NH+](C)[C@@H](C)Cc1ccccc1', 'chembl_id': ''},
    {'id': 'QFDOS-138', 'nombre': 'Rasagilina', 'nombre_en': 'Rasagiline', 'tema': 'Inhibidores MAO-B/COMT y liberadores', 'familia': 'Aminoindano propargilado', 'smiles': 'C#CCN[C@@H]1CCc2ccccc21', 'cas': '136236-51-6', 'notas': 'N-propargil-1-(R)-aminoindano. Análogo cíclico de la selegilina: la cadena fenilisopropílica queda cerrada en un indano, y el nitrógeno pasa a secundario sin metilo. Consecuencia práctica: su metabolito es el 1-aminoindano, no una anfetamina, de modo que carece del efecto estimulante. Inhibidor irreversible de la MAO-B. Configuración (R).', 'smiles_ph74': 'C#CC[NH2+][C@@H]1CCc2ccccc21', 'chembl_id': ''},
    {'id': 'QFDOS-139', 'nombre': 'Tolcapona', 'nombre_en': 'Tolcapone', 'tema': 'Inhibidores MAO-B/COMT y liberadores', 'familia': 'Nitrocatecol (benzofenona)', 'smiles': 'Cc1ccc(C(=O)c2cc(O)c(O)c([N+](=O)[O-])c2)cc1', 'cas': '134308-13-7', 'notas': 'Inhibidor reversible de la COMT, periférico y central. El nitro en orto al catecol baja el pKa del OH adyacente (≈ 4.5): a pH 7.4 circula como fenolato y es un mal sustrato de la metilación, pero se une con alta afinidad. Potencia la levodopa y permite reducir su dosis hasta un tercio. Hepatotoxicidad (restricción de uso); la entacapona, solo periférica, es la alternativa. smiles_ph74 con el fenolato orto-nitro.', 'smiles_ph74': 'Cc1ccc(C(=O)c2cc(O)c([O-])c([N+](=O)[O-])c2)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-140', 'nombre': 'Amantadina', 'nombre_en': 'Amantadine', 'tema': 'Inhibidores MAO-B/COMT y liberadores', 'familia': 'Adamantilamina', 'smiles': 'NC12CC3CC(CC(C3)C1)C2', 'cas': '768-94-5', 'notas': '1-Aminoadamantano. Único fármaco de uso clínico que provoca la liberación presináptica de dopamina (además bloquea receptores NMDA). También antivírico frente a influenza A (bloqueo del canal M2). Amina primaria sobre carbono terciario: se prepara mediante la reacción de Ritter (QFDOS-141). Jaula lipófila con TPSA de una sola NH2.', 'smiles_ph74': '[NH3+]C12CC3CC(CC(C3)C1)C2', 'chembl_id': ''},
    {'id': 'QFDOS-141', 'nombre': 'N-(1-Adamantil)acetamida', 'nombre_en': 'N-(1-Adamantyl)acetamide', 'tema': 'Inhibidores MAO-B/COMT y liberadores', 'familia': 'Intermedio de sintesis (Ritter)', 'smiles': 'CC(=O)NC12CC3CC(CC(C3)C1)C2', 'cas': '880-52-4', 'notas': 'Producto de la reacción de Ritter: el carbocatión 1-adamantilo (a partir de adamantan-1-ol o 1-bromoadamantano en medio ácido) es atrapado por el nitrilo (acetonitrilo) y el ion nitrilio se hidrata a la amida. La hidrólisis de la amida da la amantadina. Método general para aminas primarias de tipo terc-alquil-NH2, imposibles de obtener por SN2.', 'smiles_ph74': 'CC(=O)NC12CC3CC(CC(C3)C1)C2', 'chembl_id': ''},
    {'id': 'QFDOS-142', 'nombre': 'Fenotiazina', 'nombre_en': 'Phenothiazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina', 'smiles': 'c1ccc2c(c1)Nc1ccccc1S2', 'cas': '92-84-2', 'notas': 'Núcleo tricíclico (10H-dibenzo-1,4-tiazina) de los neurolépticos tricíclicos. Numeración: N10, S5, posición 2 la que recibe el grupo atrayente (Cl, CF3, SCH3...). Se obtiene por tionación de la difenilamina con azufre e I2. No es básica (NH diarilamínico).', 'smiles_ph74': 'c1ccc2c(c1)Nc1ccccc1S2', 'chembl_id': ''},
    {'id': 'QFDOS-143', 'nombre': '3-Clorodifenilamina', 'nombre_en': '3-Chlorodiphenylamine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Intermedio de sintesis', 'smiles': 'Clc1cccc(Nc2ccccc2)c1', 'cas': '101-17-6', 'notas': 'Sustrato de la tionación (S8, I2, calor) para la síntesis de la 2-clorofenotiazina. El cierre puede producirse en orto o para al cloro y genera dos isómeros (2-cloro y 4-cloro) que deben separarse; esta falta de regioselectividad justifica las rutas alternativas de Ullmann y transposición de Smiles.', 'smiles_ph74': 'Clc1cccc(Nc2ccccc2)c1', 'chembl_id': ''},
    {'id': 'QFDOS-144', 'nombre': '2-Clorofenotiazina', 'nombre_en': '2-Chlorophenothiazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina', 'smiles': 'Clc1ccc2Sc3ccccc3Nc2c1', 'cas': '92-39-7', 'notas': 'Intermedio clave de la clorpromazina. Se N-alquila con NaNH2 y 3-cloro-N,N-dimetilpropilamina. El Cl en 2 es el grupo atrayente que optimiza la actividad neuroléptica (región C del farmacóforo de Gordon).', 'smiles_ph74': 'Clc1ccc2c(c1)Nc1ccccc1S2', 'chembl_id': ''},
    {'id': 'QFDOS-145', 'nombre': 'Prometazina', 'nombre_en': 'Promethazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina', 'smiles': 'CC(CN1c2ccccc2Sc2ccccc21)N(C)C', 'cas': '60-87-7', 'notas': 'Fenotiazina antihistamínica H1 con efecto sedante: puente de 2 carbonos ramificado con metilo entre los nitrógenos y sin sustituyente en 2. Ilustra dos reglas del farmacóforo: con 2 carbonos y ramificación predomina el antihistamínico sobre el neuroléptico. Puente histórico entre las etilendiaminas antihistamínicas y la clorpromazina. Racémico; se guarda sin estereoquímica.', 'smiles_ph74': 'CC(CN1c2ccccc2Sc2ccccc21)[NH+](C)C', 'chembl_id': ''},
    {'id': 'QFDOS-146', 'nombre': 'Clorpromazina', 'nombre_en': 'Chlorpromazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina', 'smiles': 'CN(C)CCCN1c2ccccc2Sc2ccc(Cl)cc21', 'cas': '50-53-3', 'notas': 'Prototipo de los neurolépticos (1952). Cumple las tres zonas del farmacóforo de Gordon: A, amina terciaria protonable; B, cadena de exactamente 3 carbonos entre nitrógenos; C, tricíclico con Cl en 2. Antagonista D2 con acciones sobre H1, alfa1 y muscarínicos. La conformación preferida dispone la cadena hacia el anillo clorado, superponible con la dopamina antiperiplanar. Sintetizable por tionación, por Ullmann o por transposición de Smiles.', 'smiles_ph74': 'C[NH+](C)CCCN1c2ccccc2Sc2ccc(Cl)cc21', 'chembl_id': ''},
    {'id': 'QFDOS-147', 'nombre': 'Levomepromazina', 'nombre_en': 'Levomepromazine (methotrimeprazine)', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina', 'smiles': 'COc1ccc2Sc3ccccc3N(C[C@H](C)CN(C)C)c2c1', 'cas': '60-99-1', 'notas': 'Metotrimeprazina (Sinogán). 2-Metoxi y cadena de 3 carbonos ramificada con metilo en β: perfil intermedio entre clorpromazina y prometazina, muy sedante y analgésico. Eutómero (R)-(−) (levo). Ejemplo de que el OMe, menos atrayente que Cl o CF3, rebaja la potencia antipsicótica a favor de la sedación.', 'smiles_ph74': 'COc1ccc2c(c1)N(C[C@H](C)C[NH+](C)C)c1ccccc1S2', 'chembl_id': ''},
    {'id': 'QFDOS-148', 'nombre': 'Carfenazina', 'nombre_en': 'Carphenazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina (piperazinil)', 'smiles': 'CCC(=O)c1ccc2Sc3ccccc3N(CCCN3CCN(CCO)CC3)c2c1', 'cas': '2622-30-2', 'notas': '2-Propionilfenotiazina con cadena hidroxietilpiperazinilpropilo. Síntesis: N-acilación protectora con EtCOCl, Friedel-Crafts en 2 (el S dirige al desaparecer la activación del N acilado), desprotección, alquilación con 1-bromo-3-cloropropano (reacciona el Br) y sustitución por 1-(2-hidroxietil)piperazina (QFDOS-152). El OH terminal permite esterificar para depot, como en la flufenazina.', 'smiles_ph74': 'CCC(=O)c1ccc2c(c1)N(CCCN1CC[NH+](CCO)CC1)c1ccccc1S2', 'chembl_id': ''},
    {'id': 'QFDOS-149', 'nombre': 'Flufenazina', 'nombre_en': 'Fluphenazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina (piperazinil)', 'smiles': 'OCCN1CCN(CCCN2c3ccccc3Sc3ccc(C(F)(F)F)cc32)CC1', 'cas': '69-23-8', 'notas': 'Neuroléptico piperazínico de alta potencia. CF3 en 2 (más atrayente y lipófilo que Cl) y N básico incluido en una piperazina, dos modificaciones que aumentan la potencia D2 frente a la clorpromazina y desplazan el perfil a más efectos extrapiramidales y menos sedación. El hidroxietilo terminal es el punto de anclaje de los ésteres depot (QFDOS-150, 151).', 'smiles_ph74': 'OCC[NH+]1CCN(CCCN2c3ccccc3Sc3ccc(C(F)(F)F)cc32)CC1', 'chembl_id': ''},
    {'id': 'QFDOS-150', 'nombre': 'Decanoato de flufenazina', 'nombre_en': 'Fluphenazine decanoate', 'tema': 'Neurolepticos triciclicos', 'familia': 'Profarmaco depot (ester)', 'smiles': 'CCCCCCCCCC(=O)OCCN1CCN(CCCN2c3ccccc3Sc3ccc(C(F)(F)F)cc32)CC1', 'cas': '5002-47-1', 'notas': 'Éster del ácido decanoico (cáprico) sobre el OH de la flufenazina. Latentización para acción muy prolongada: en solución oleosa intramuscular se libera lentamente del depósito y se hidroliza a flufenazina, con efecto de 2 a 4 semanas. Útil cuando la adherencia oral es el problema. logP muy alto: es el éster, no la especie activa, lo que describe el descriptor.', 'smiles_ph74': 'CCCCCCCCCC(=O)OCC[NH+]1CCN(CCCN2c3ccccc3Sc3ccc(C(F)(F)F)cc32)CC1', 'chembl_id': ''},
    {'id': 'QFDOS-151', 'nombre': 'Enantato de flufenazina', 'nombre_en': 'Fluphenazine enanthate', 'tema': 'Neurolepticos triciclicos', 'familia': 'Profarmaco depot (ester)', 'smiles': 'CCCCCCC(=O)OCCN1CCN(CCCN2c3ccccc3Sc3ccc(C(F)(F)F)cc32)CC1', 'cas': '2746-81-8', 'notas': 'Éster del ácido heptanoico (enántico). Cadena tres carbonos más corta que el decanoato: liberación algo más rápida y duración menor. El par enantato/decanoato muestra cómo la longitud de cadena del éster regula la velocidad de salida del depósito.', 'smiles_ph74': 'CCCCCCC(=O)OCC[NH+]1CCN(CCCN2c3ccccc3Sc3ccc(C(F)(F)F)cc32)CC1', 'chembl_id': ''},
    {'id': 'QFDOS-152', 'nombre': '1-(2-Hidroxietil)piperazina', 'nombre_en': '1-(2-Hydroxyethyl)piperazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Intermedio de sintesis', 'smiles': 'OCCN1CCNCC1', 'cas': '103-76-4', 'notas': 'Reactivo para introducir la cadena lateral de carfenazina y flufenazina. Se prepara protegiendo un N de la piperazina como uretano (ClCOOEt), alquilando el otro con óxido de etileno y desprotegiendo con NaOH: monofuncionalización selectiva de una diamina simétrica.', 'smiles_ph74': 'OCCN1CC[NH2+]CC1', 'chembl_id': ''},
    {'id': 'QFDOS-153', 'nombre': 'Clorprotixeno', 'nombre_en': 'Chlorprothixene', 'tema': 'Neurolepticos triciclicos', 'familia': 'Tioxanteno', 'smiles': 'CN(C)CC/C=C1/c2ccccc2Sc2ccc(Cl)cc21', 'cas': '113-59-7', 'notas': 'Tioxanteno: el N10 de la fenotiazina se sustituye por un carbono sp2 unido a la cadena por un doble enlace exocíclico. Isosterismo N→C que conserva el farmacóforo. El doble enlace genera isómeros geométricos; el isómero (Z), con la cadena hacia el anillo clorado, es el activo y el (E) mucho menos, lo que confirma la conformación preferida del farmacóforo de Gordon. Verificar la geometría con una fuente independiente.', 'smiles_ph74': 'C[NH+](C)CC/C=C1/c2ccccc2Sc2ccc(Cl)cc21', 'chembl_id': ''},
    {'id': 'QFDOS-154', 'nombre': 'Estructura 1 (Ejercicio 3.3)', 'nombre_en': 'Exercise 3.3 rigid phenothiazine', 'tema': 'Neurolepticos triciclicos', 'familia': 'Fenotiazina rigida', 'smiles': 'CN(C)C1CN2c3ccccc3Sc3cccc(C1)c32', 'cas': '', 'notas': 'Fenotiazina tetracíclica del ejercicio 3.3: la cadena trimetilénica se cierra sobre el carbono peri de un anillo bencénico. Tiene amina terciaria, 3 carbonos entre nitrógenos y tricíclico, pero es inactiva: la rigidez impide la conformación extendida de la cadena y por tanto la superposición con la dopamina antiperiplanar. Contraejemplo del farmacóforo. Estructura dibujada en la diapositiva, sin referencia en bases de datos; racémica.', 'smiles_ph74': 'C[NH+](C)C1Cc2cccc3c2N(C1)c1ccccc1S3', 'chembl_id': ''},
    {'id': 'QFDOS-155', 'nombre': 'Petidina', 'nombre_en': 'Pethidine (meperidine)', 'tema': 'Butirofenonas y analogos', 'familia': '4-Fenilpiperidina', 'smiles': 'CCOC(=O)C1(c2ccccc2)CCN(C)CC1', 'cas': '57-42-1', 'notas': 'Analgésico opioide sintético (Dolantina), simplificación de la morfina por supresión de los anillos B, C y D: conserva fenilo, carbono cuaternario y piperidina N-metilada. Punto de partida de las butirofenonas: sustituyendo el N-metilo por una cadena butirofenónica se obtiene un compuesto analgésico y neuroléptico (QFDOS-156).', 'smiles_ph74': 'CCOC(=O)C1(c2ccccc2)CC[NH+](C)CC1', 'chembl_id': ''},
    {'id': 'QFDOS-156', 'nombre': 'Analogo butirofenonico de la petidina', 'nombre_en': 'Pethidine butyrophenone analogue', 'tema': 'Butirofenonas y analogos', 'familia': 'Butirofenona', 'smiles': 'CCOC(=O)C1(c2ccccc2)CCN(CCCC(=O)c2ccccc2)CC1', 'cas': '', 'notas': 'Intermedio conceptual (analgésico y neuroléptico). La cadena 4-oxo-4-fenilbutilo sobre el N de la petidina aporta la acción neuroléptica. Introducir F en para del anillo acilado y cambiar el éster por OH separa ambas acciones y conduce al haloperidol.', 'smiles_ph74': 'CCOC(=O)C1(c2ccccc2)CC[NH+](CCCC(=O)c2ccccc2)CC1', 'chembl_id': ''},
    {'id': 'QFDOS-157', 'nombre': 'Haloperidol', 'nombre_en': 'Haloperidol', 'tema': 'Butirofenonas y analogos', 'familia': 'Butirofenona', 'smiles': 'O=C(CCCN1CCC(O)(c2ccc(Cl)cc2)CC1)c1ccc(F)cc1', 'cas': '52-86-8', 'notas': "Prototipo de las butirofenonas: 4-[4-(4-clorofenil)-4-hidroxipiperidino]-4'-fluorobutirofenona. Antagonista D2 de alta potencia y escasa acción sobre otros receptores: poco sedante e hipotensor, muchos efectos extrapiramidales. Farmacóforo: p-fluorobutirofenona + N básico a 3 carbonos del carbonilo + 4-aril-4-piperidinol. Síntesis por la vía de la tetrahidropiridina (reacción tipo Mannich/Prins sobre α-metil-p-cloroestireno). Decanoato depot.", 'smiles_ph74': 'O=C(CCC[NH+]1CCC(O)(c2ccc(Cl)cc2)CC1)c1ccc(F)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-158', 'nombre': '4-(4-Clorofenil)piperidin-4-ol', 'nombre_en': '4-(4-Chlorophenyl)piperidin-4-ol', 'tema': 'Butirofenonas y analogos', 'familia': 'Intermedio de sintesis', 'smiles': 'OC1(c2ccc(Cl)cc2)CCNCC1', 'cas': '39512-49-7', 'notas': "Fragmento 4-aril-4-piperidinol del haloperidol. En la síntesis: tetrahidropiridina (QFDOS-161) + HBr/AcOH (adición Markovnikov del acetato o bromuro) e hidrólisis con NaOH. Se N-alquila con 4-cloro-4'-fluorobutirofenona (Na2CO3, KI, n-butanol; el KI activa por Finkelstein).", 'smiles_ph74': 'OC1(c2ccc(Cl)cc2)CC[NH2+]CC1', 'chembl_id': ''},
    {'id': 'QFDOS-159', 'nombre': "4-Cloro-4'-fluorobutirofenona", 'nombre_en': "4-Chloro-4'-fluorobutyrophenone", 'tema': 'Butirofenonas y analogos', 'familia': 'Intermedio de sintesis', 'smiles': 'O=C(CCCCl)c1ccc(F)cc1', 'cas': '3874-54-2', 'notas': 'Agente alquilante común a casi todas las butirofenonas (haloperidol, droperidol, trifluperidol). Se obtiene por Friedel-Crafts del fluorobenceno con cloruro de 4-clorobutanoílo.', 'smiles_ph74': 'O=C(CCCCl)c1ccc(F)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-160', 'nombre': 'alfa-Metil-p-cloroestireno', 'nombre_en': '4-Chloro-alpha-methylstyrene', 'tema': 'Butirofenonas y analogos', 'familia': 'Intermedio de sintesis', 'smiles': 'C=C(C)c1ccc(Cl)cc1', 'cas': '1712-70-5', 'notas': 'Alqueno de partida de la síntesis del haloperidol. Reacciona con formaldehído y NH4Cl (catión iminio CH2=NH2+): adición electrófila al alqueno que recuerda a la reacción de Mannich, segunda adición de formaldehído y ciclación a una 1,3-oxazinana, que en HCl se transforma en la tetrahidropiridina.', 'smiles_ph74': 'C=C(C)c1ccc(Cl)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-161', 'nombre': '4-(4-Clorofenil)-1,2,3,6-tetrahidropiridina', 'nombre_en': '4-(4-Chlorophenyl)-1,2,3,6-tetrahydropyridine', 'tema': 'Butirofenonas y analogos', 'familia': 'Intermedio de sintesis', 'smiles': 'Clc1ccc(C2=CCNCC2)cc1', 'cas': '30005-58-4', 'notas': 'Intermedio del haloperidol (vía Mannich/Prins). Mismo tipo de anillo que el droperidol conserva en el fármaco final. Estructuralmente relacionada con la MPTP (1-metil-4-fenil-1,2,3,6-tetrahidropiridina), neurotoxina que oxidada por la MAO-B a MPP+ destruye neuronas dopaminérgicas y produce parkinsonismo, el modelo experimental del Parkinson.', 'smiles_ph74': 'Clc1ccc(C2=CC[NH2+]CC2)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-162', 'nombre': 'Droperidol', 'nombre_en': 'Droperidol', 'tema': 'Butirofenonas y analogos', 'familia': 'Butirofenona', 'smiles': 'O=C(CCCN1CC=C(n2c(=O)[nH]c3ccccc32)CC1)c1ccc(F)cc1', 'cas': '548-73-2', 'notas': 'Butirofenona en la que el 4-arilpiperidinol se sustituye por una 4-(2-oxobencimidazolinil)-1,2,3,6-tetrahidropiridina. Neuroléptico de acción corta y antiemético; en anestesia (neuroleptoanalgesia con fentanilo). Riesgo de prolongación del QT.', 'smiles_ph74': 'O=C(CCC[NH+]1CC=C(n2c(=O)[nH]c3ccccc32)CC1)c1ccc(F)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-163', 'nombre': 'Pimozida', 'nombre_en': 'Pimozide', 'tema': 'Butirofenonas y analogos', 'familia': 'Difenilbutilpiperidina', 'smiles': 'O=c1[nH]c2ccccc2n1C1CCN(CCCC(c2ccc(F)cc2)c2ccc(F)cc2)CC1', 'cas': '2062-78-4', 'notas': 'Difenilbutilpiperidina: el carbonilo de la butirofenona se sustituye por un segundo anillo p-fluorofenilo, demostrando que la cetona no es imprescindible. Conserva el bencimidazolona-piperidina del droperidol. Muy lipófila y de semivida larga (una dosis diaria). Bloqueo de hERG: riesgo de QT largo.', 'smiles_ph74': 'O=c1[nH]c2ccccc2n1C1CC[NH+](CCCC(c2ccc(F)cc2)c2ccc(F)cc2)CC1', 'chembl_id': ''},
    {'id': 'QFDOS-164', 'nombre': 'Trifluperidol', 'nombre_en': 'Trifluperidol', 'tema': 'Butirofenonas y analogos', 'familia': 'Butirofenona', 'smiles': 'O=C(CCCN1CCC(O)(c2cccc(C(F)(F)F)c2)CC1)c1ccc(F)cc1', 'cas': '749-13-3', 'notas': 'Análogo del haloperidol con 3-CF3-fenilo en lugar de 4-clorofenilo en el piperidinol. Más potente. Objeto del ejercicio 3.6 (completar el esquema de síntesis): misma secuencia que el haloperidol partiendo del alqueno o del organometálico con 3-trifluorometilfenilo.', 'smiles_ph74': 'O=C(CCC[NH+]1CCC(O)(c2cccc(C(F)(F)F)c2)CC1)c1ccc(F)cc1', 'chembl_id': ''},
    {'id': 'QFDOS-165', 'nombre': 'orto-Metoxiprocainamida', 'nombre_en': 'ortho-Methoxyprocainamide', 'tema': 'Ortopramidas', 'familia': 'Benzamida', 'smiles': 'CCN(CC)CCNC(=O)c1ccc(N)cc1OC', 'cas': '', 'notas': 'Procainamida con 2-OMe. Mostró, además de la acción anestésica local esperada, una notable actividad antiemética. Cabeza de serie de las ortopramidas. El OMe en orto forma un enlace de hidrógeno intramolecular con el NH amídico que fija un pseudoanillo coplanar, conformación clave para la unión a D2.', 'smiles_ph74': 'CC[NH+](CC)CCNC(=O)c1ccc(N)cc1OC', 'chembl_id': ''},
    {'id': 'QFDOS-166', 'nombre': 'Metoclopramida', 'nombre_en': 'Metoclopramide', 'tema': 'Ortopramidas', 'familia': 'Benzamida (ortopramida)', 'smiles': 'CCN(CC)CCNC(=O)c1cc(Cl)c(N)cc1OC', 'cas': '364-62-5', 'notas': '4-Amino-5-cloro-2-metoxi-N-(2-dietilaminoetil)benzamida (Primperan). Antagonista D2 en la zona quimiorreceptora gatillo: antiemético potente, procinético (también agonista 5-HT4), anestésico local moderado. Cruza la BHE: efectos extrapiramidales. Respecto a la orto-metoxiprocainamida añade solo un Cl en 5.', 'smiles_ph74': 'CC[NH+](CC)CCNC(=O)c1cc(Cl)c(N)cc1OC', 'chembl_id': ''},
    {'id': 'QFDOS-167', 'nombre': 'Dopamina', 'nombre_en': 'Dopamine', 'tema': 'Dopamina: farmaco y diana', 'familia': 'Catecolamina', 'smiles': 'NCCc1ccc(O)c(O)c1', 'smiles_ph74': '[NH3+]CCc1ccc(O)c(O)c1', 'cas': '51-61-6', 'chembl_id': 'CHEMBL59', 'notas': 'Protagonista del Tema 3 (misma molécula que QFDOS-008 y QFDOS-048, aquí con su contexto dopaminérgico). Catecol + cadena etilamínica de libre giro: conformaciones múltiples, de las que la antiperiplanar es la activa (argumento de ADTN y apomorfina, QFDOS-128/129). Amina primaria protonada a pH 7.4 y catecol muy polar: no se absorbe bien por vía oral ni cruza la BHE, de ahí la levodopa (QFDOS-168) y los profármacos latentizados (QFDOS-132/133). Degradada por MAO-B en el estriado y por COMT (dianas de selegilina, rasagilina y tolcapona). Su déficit en la sustancia negra es la base bioquímica del Parkinson; su exceso en la vía mesolímbica, la diana de los neurolépticos D2. Agonista de la zona quimiorreceptora gatillo: emética.'},
    {'id': 'QFDOS-168', 'nombre': 'L-DOPA (Levodopa)', 'nombre_en': 'Levodopa', 'tema': 'Dopamina: farmaco y diana', 'familia': 'Catecolaminoacido', 'smiles': 'N[C@@H](Cc1ccc(O)c(O)c1)C(=O)O', 'smiles_ph74': '[NH3+][C@@H](Cc1ccc(O)c(O)c1)C(=O)[O-]', 'cas': '59-92-7', 'chembl_id': 'CHEMBL1021', 'notas': 'Misma molécula que QFDOS-047, aquí como antiparkinsoniano. Profármaco de la dopamina: el resto aminoácido la hace sustrato del transportador LAT1, que la lleva a través de la mucosa intestinal y de la BHE; dentro del cerebro la AADC la descarboxila a dopamina. Se asocia siempre a un inhibidor periférico de la AADC (carbidopa QFDOS-169, benserazida QFDOS-131) para evitar la dopamina periférica (náuseas, arritmias) y reducir la dosis, y a menudo a un inhibidor de COMT (tolcapona) o MAO-B (selegilina, rasagilina). Configuración (S).'},
    {'id': 'QFDOS-169', 'nombre': 'Carbidopa', 'nombre_en': 'Carbidopa', 'tema': 'Dopamina: farmaco y diana', 'familia': 'Catecolhidrazina', 'smiles': 'C[C@](NN)(Cc1ccc(O)c(O)c1)C(=O)O', 'smiles_ph74': 'C[C@](N[NH3+])(Cc1ccc(O)c(O)c1)C(=O)[O-]', 'cas': '28860-95-9', 'chembl_id': 'CHEMBL1201236', 'notas': 'Misma molécula que QFDOS-094. Inhibidor de la dopa-descarboxilasa periférica (Sinemet con levodopa): análogo α-hidrazínico de la α-metildopa cuya hidrazina forma una hidrazona con el piridoxal fosfato de la AADC. Zwitterión polar que no cruza la BHE, de modo que la descarboxilación central de la levodopa queda intacta. Par con la benserazida (QFDOS-131): mismo mecanismo (hidrazina frente a PLP) sobre esqueletos distintos. Configuración (S).'}]

# ─── 2. PIPELINE DE ANÁLISIS RDKIT Y DESCRIPTORES ─────────────────────────────

def descriptores_cip(mol):
    Chem.AssignStereochemistry(mol, cleanIt=True, force=True)
    from rdkit.Chem import rdCIPLabeler
    try:
        rdCIPLabeler.AssignCIPLabels(mol)
    except Exception:
        pass
    etiquetas = []
    for a in mol.GetAtoms():
        if a.HasProp('_CIPCode'):
            etiquetas.append('C%d:%s' % (a.GetIdx() + 1, a.GetProp('_CIPCode')))
    sin_asignar = [i for i, c in Chem.FindMolChiralCenters(
        mol, includeUnassigned=True, useLegacyImplementation=False)
        if c == '?' and mol.GetAtomWithIdx(i).GetSymbol() in ('C', 'N')]
    if etiquetas:
        texto = '; '.join(etiquetas)
        if sin_asignar:
            texto += ' | %d centro(s) sin definir' % len(sin_asignar)
    elif sin_asignar:
        texto = '%d centro(s) sin definir' % len(sin_asignar)
    else:
        texto = 'aquiral'
    return texto


def analizar(smiles):
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError('SMILES no valido: %s' % smiles)
    d = {}
    d['smiles_canonico'] = Chem.MolToSmiles(mol)
    d['formula'] = rdMolDescriptors.CalcMolFormula(mol)
    d['mw'] = round(Descriptors.MolWt(mol), 2)
    d['masa_exacta'] = round(Descriptors.ExactMolWt(mol), 4)
    d['carga_formal'] = Chem.GetFormalCharge(mol)
    d['logp'] = round(Crippen.MolLogP(mol), 2)
    d['refractividad_molar'] = round(Crippen.MolMR(mol), 2)
    d['tpsa'] = round(Descriptors.TPSA(mol), 2)
    d['hbd'] = Lipinski.NumHDonors(mol)
    d['hba'] = Lipinski.NumHAcceptors(mol)
    d['enlaces_rotables'] = rdMolDescriptors.CalcNumRotatableBonds(mol)
    d['atomos_pesados'] = mol.GetNumHeavyAtoms()
    d['anillos'] = rdMolDescriptors.CalcNumRings(mol)
    d['anillos_aromaticos'] = rdMolDescriptors.CalcNumAromaticRings(mol)
    d['fraccion_csp3'] = round(rdMolDescriptors.CalcFractionCSP3(mol), 3)
    d['qed'] = round(QED.qed(mol), 3)
    d['estereocentros_cip'] = descriptores_cip(mol)
    d['inchi'] = Chem.MolToInchi(mol)
    d['inchikey'] = Chem.MolToInchiKey(mol)
    viol = sum([d['mw'] > 500, d['logp'] > 5, d['hbd'] > 5, d['hba'] > 10])
    d['lipinski_violaciones'] = viol
    d['cumple_lipinski'] = 'Si' if viol <= 1 else 'No'
    return mol, d


def leyenda_segura(texto):
    """Pliega los acentos a ASCII para las leyendas de las imágenes.

    La fuente por omisión de RDKit no incorpora los glifos acentuados y, en lugar de
    dibujarlos, los descarta: 'Ácido acetilsalicílico' se rotulaba como 'cido
    acetilsalic lico'. Plegar a ASCII conserva la palabra legible. Solo afecta al
    rótulo de la imagen; el CSV, el XLSX y las notas docentes mantienen los acentos.
    """
    import unicodedata
    normalizado = unicodedata.normalize('NFKD', texto)
    return ''.join(c for c in normalizado if not unicodedata.combining(c))


def dibujar(mol, ruta, leyenda='', ancho=500, alto=400):
    m = Chem.Mol(mol)
    AllChem.Compute2DCoords(m)
    drawer = rdMolDraw2D.MolDraw2DCairo(ancho, alto)
    opts = drawer.drawOptions()
    opts.clearBackground = True
    opts.addStereoAnnotation = True
    opts.bondLineWidth = 2.0
    opts.padding = 0.05
    rdMolDraw2D.PrepareAndDrawMolecule(drawer, m, legend=leyenda_segura(leyenda))
    drawer.FinishDrawing()
    with open(ruta, 'wb') as fh:
        fh.write(drawer.GetDrawingText())


def slug(texto):
    t = texto.lower()
    t = (t.replace('á', 'a').replace('é', 'e').replace('í', 'i')
          .replace('ó', 'o').replace('ú', 'u').replace('ñ', 'n')
          .replace('α', 'a').replace('β', 'b'))
    t = re.sub(r'[^a-z0-9]+', '_', t).strip('_')
    return t


def exportar_xlsx(salida, ruta_xlsx):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Propiedades'
    cols = [c for c in salida[0].keys() if c not in ('inchi',)]
    ws.append(cols)
    cab = PatternFill('solid', fgColor='1F3864')
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = cab
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in salida:
        ws.append([r[c] for c in cols])
    anchos = {'nombre': 28, 'nombre_en': 22, 'familia': 22, 'tema': 26,
              'smiles_entrada': 34, 'smiles_ph74': 34, 'smiles_canonico': 34,
              'inchikey': 30, 'notas': 60, 'estereocentros_cip': 18, 'imagen': 30}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(c, 14)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet('Estructuras')
    ws2.append(['id', 'nombre', 'formula', 'estructura'])
    for i, c in enumerate(['id', 'nombre', 'formula', 'estructura'], 1):
        ws2.cell(row=1, column=i).font = Font(bold=True, color='FFFFFF')
        ws2.cell(row=1, column=i).fill = cab
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 30
    ws2.column_dimensions['C'].width = 16
    ws2.column_dimensions['D'].width = 42
    fila = 2
    for r in salida:
        ws2.cell(row=fila, column=1, value=r['id'])
        ws2.cell(row=fila, column=2, value=r['nombre'])
        ws2.cell(row=fila, column=3, value=r['formula'])
        ruta = os.path.join(TEMA3_DIR, r['imagen'])
        if os.path.exists(ruta):
            img = XLImage(ruta)
            img.width, img.height = 250, 200
            ws2.row_dimensions[fila].height = 155
            ws2.add_image(img, 'D%d' % fila)
        fila += 1

    wb.save(ruta_xlsx)
    print('-> %s' % ruta_xlsx)



def generar_figuras_compuestas():
    print('\nGenerando figuras didacticas compuestas...')
    extra = {'QFDOS-068':('Metanfetamina','CN[C@@H](C)Cc1ccccc1')}
    mols = {r['id']: Chem.MolFromSmiles(r['smiles']) for r in COMPUESTOS_TEMA3}
    names = {r['id']: r['nombre'] for r in COMPUESTOS_TEMA3}
    for k,(n,s) in extra.items():
        mols[k]=Chem.MolFromSmiles(s); names[k]=n
    def grid_img(ids, out_name, mols_per_row=3, sub_w=400, sub_h=300):
        img = Draw.MolsToGridImage([mols[i] for i in ids], molsPerRow=mols_per_row, subImgSize=(sub_w, sub_h),
                                   legends=[leyenda_segura(f"{i}\n{names[i]}") for i in ids], returnPNG=False)
        img.save(os.path.join(TEMA3_DIR, out_name)); print('  ->', out_name)
    grid_img(['QFDOS-167','QFDOS-127','QFDOS-128','QFDOS-129','QFDOS-130'], 'QFDOS_T3_bloque1_analogos_restringidos_dopamina.png', 5)
    grid_img(['QFDOS-168','QFDOS-169','QFDOS-131','QFDOS-132'], 'QFDOS_T3_bloque2_levodopa_inhibidores_AADC_dipivaloil.png', 4)
    grid_img(['QFDOS-133','QFDOS-134','QFDOS-167','QFDOS-135'], 'QFDOS_T3_inversion_polaridad_doble_profarmaco.png', 4, 450, 330)
    grid_img(['QFDOS-136','QFDOS-137','QFDOS-068','QFDOS-138','QFDOS-139','QFDOS-140'], 'QFDOS_T3_bloque3_MAOB_COMT_amantadina.png', 3)
    grid_img(['QFDOS-142','QFDOS-145','QFDOS-146','QFDOS-147','QFDOS-149','QFDOS-153'], 'QFDOS_T3_bloque4_fenotiazinas_tioxantenos.png', 3)
    grid_img(['QFDOS-143','QFDOS-144','QFDOS-146'], 'QFDOS_T3_sintesis_clorpromazina.png', 3)
    grid_img(['QFDOS-148','QFDOS-149','QFDOS-151','QFDOS-150'], 'QFDOS_T3_esteres_depot_flufenazina.png', 2, 500, 350)
    grid_img(['QFDOS-146','QFDOS-154'], 'QFDOS_T3_ejercicio_3_3_fenotiazina_rigida.png', 2, 450, 330)
    grid_img(['QFDOS-130','QFDOS-155','QFDOS-156','QFDOS-157'], 'QFDOS_T3_de_petidina_a_haloperidol.png', 4, 420, 320)
    grid_img(['QFDOS-160','QFDOS-161','QFDOS-158','QFDOS-159','QFDOS-157'], 'QFDOS_T3_sintesis_haloperidol.png', 5)
    grid_img(['QFDOS-157','QFDOS-164','QFDOS-162','QFDOS-163'], 'QFDOS_T3_bloque5_butirofenonas_difenilbutilpiperidinas.png', 4, 420, 320)
    grid_img(['QFDOS-165','QFDOS-166'], 'QFDOS_T3_bloque6_ortopramidas.png', 2, 450, 330)

def main():
    print(f'=== PROCESANDO TEMA 3 QFDOS ({len(COMPUESTOS_TEMA3)} ESTRUCTURAS) ===\n')

    # Guardar CSV maestro inicial en Tema 3
    maestro_csv = os.path.join(TEMA3_DIR, 'estructuras_qfdos.csv')
    cabeceras_maestro = ['id', 'nombre', 'nombre_en', 'tema', 'familia', 'smiles',
                         'smiles_ph74', 'cas', 'chembl_id', 'notas']
    with open(maestro_csv, 'w', newline='', encoding='utf-8') as fh:
        w = csv.DictWriter(fh, fieldnames=cabeceras_maestro)
        w.writeheader()
        w.writerows(COMPUESTOS_TEMA3)
    print(f'-> {maestro_csv} guardado con éxito.')

    salida = []
    for r in COMPUESTOS_TEMA3:
        mol, d = analizar(r['smiles'])

        # Rutas de imagen: en img/ y en raíz de Tema 3
        slug_nom = slug(r['nombre'])
        img_rel = 'img/%s_%s.png' % (r['id'], slug_nom)
        img_abs = os.path.join(TEMA3_DIR, img_rel)
        img_root = os.path.join(TEMA3_DIR, '%s_%s.png' % (r['id'], slug_nom))

        dibujar(mol, img_abs, leyenda=r['nombre'])
        dibujar(mol, img_root, leyenda=r['nombre'])

        carga_ph74 = ''
        if r.get('smiles_ph74'):
            m74 = Chem.MolFromSmiles(r['smiles_ph74'])
            if m74 is not None:
                carga_ph74 = Chem.GetFormalCharge(m74)

        fila = {
            'id': r['id'], 'nombre': r['nombre'], 'nombre_en': r['nombre_en'],
            'tema': r['tema'], 'familia': r['familia'], 'cas': r['cas'],
            'chembl_id': r.get('chembl_id', ''),
            'smiles_entrada': r['smiles'], 'smiles_ph74': r.get('smiles_ph74', ''),
            'carga_neta_ph74': carga_ph74,
        }
        fila.update(d)
        fila['imagen'] = img_rel
        fila['notas'] = r['notas']
        salida.append(fila)

        print('[OK] %-12s %-28s %-12s MW %7.2f  logP %6.2f  TPSA %6.2f  %s'
              % (r['id'], r['nombre'][:28], d['formula'], d['mw'], d['logp'],
                 d['tpsa'], d['inchikey']))

    # Exportar propiedades CSV
    propiedades_csv = os.path.join(TEMA3_DIR, 'propiedades_qfdos.csv')
    cols = list(salida[0].keys())
    with open(propiedades_csv, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(salida)
    print(f'\n-> {propiedades_csv} ({len(salida)} estructuras)')

    # Exportar Excel XLSX
    try:
        ruta_xlsx = os.path.join(TEMA3_DIR, 'estructuras_qfdos.xlsx')
        exportar_xlsx(salida, ruta_xlsx)
    except Exception as e:
        print(f'[aviso] Error al exportar XLSX: {e}')

    # Generar figuras didácticas compuestas
    generar_figuras_compuestas()

    print('\n¡PROCESO COMPLETADO EXITOSAMENTE PARA EL TEMA 3!')


if __name__ == '__main__':
    main()
