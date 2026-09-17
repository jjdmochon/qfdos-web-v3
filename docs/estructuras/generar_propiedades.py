# -*- coding: utf-8 -*-
"""
QFDOS - Pipeline de estructuras quimicas del curso
==================================================
Lee el CSV maestro (estructuras_qfdos.csv), valida los SMILES, calcula
descriptores fisicoquimicos con RDKit y genera:

  - propiedades_qfdos.csv   : tabla de descriptores
  - estructuras_qfdos.xlsx  : libro Excel (maestro + propiedades + imagenes)
  - img/<id>_<nombre>.png   : depiccion 2D de cada estructura

Uso:
    python generar_propiedades.py                 # usa ./estructuras_qfdos.csv
    python generar_propiedades.py otro_maestro.csv

Para anadir estructuras basta editar el CSV maestro y volver a ejecutar.
Requisitos: rdkit, pandas, openpyxl
"""
import os
import sys
import csv

from rdkit import Chem, RDLogger
from rdkit.Chem import (AllChem, Crippen, Descriptors, Draw, Lipinski,
                        rdMolDescriptors)
from rdkit.Chem.Draw import rdMolDraw2D
from rdkit.Chem import QED

RDLogger.DisableLog('rdApp.*')

BASE = os.path.dirname(os.path.abspath(__file__))
MAESTRO = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, 'estructuras_qfdos.csv')
IMGDIR = os.path.join(BASE, 'img')
os.makedirs(IMGDIR, exist_ok=True)

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass


def descriptores_cip(mol):
    """Devuelve los descriptores estereoquimicos CIP, p.ej. 'C2:S'."""
    Chem.AssignStereochemistry(mol, cleanIt=True, force=True)
    from rdkit.Chem import rdCIPLabeler
    try:
        rdCIPLabeler.AssignCIPLabels(mol)
    except Exception:
        pass
    etiquetas = []
    for a in mol.GetAtoms():
        if a.HasProp('_CIPCode'):
            etiquetas.append('C%d:%s' % (a.GetIdx() + 1, a.GetProp('_CIPCode')))
    # Centros presentes pero sin configuracion definida: racematos (carbono) o
    # nitrogenos estereogenicos que la fuente no resuelve (N-oxidos, amonios cuaternarios)
    # Se excluye el fosforo: RDKit marca como potencialmente estereogenicos los
    # P de fosfatos y pirofosfatos (=O / -OH intercambiables por ionizacion y
    # tautomeria), que no son centros reales. Solo cuentan carbono y nitrogeno.
    sin_asignar = [i for i, c in Chem.FindMolChiralCenters(
        mol, includeUnassigned=True, useLegacyImplementation=False)
        if c == '?' and mol.GetAtomWithIdx(i).GetSymbol() in ('C', 'N')]
    if etiquetas:
        texto = '; '.join(etiquetas)
        if sin_asignar:
            texto += ' | %d centro(s) sin definir' % len(sin_asignar)
    elif sin_asignar:
        texto = '%d centro(s) sin definir' % len(sin_asignar)
    else:
        texto = 'aquiral'
    return texto


def analizar(smiles):
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError('SMILES no valido: %s' % smiles)
    d = {}
    d['smiles_canonico'] = Chem.MolToSmiles(mol)
    d['formula'] = rdMolDescriptors.CalcMolFormula(mol)
    d['mw'] = round(Descriptors.MolWt(mol), 2)
    d['masa_exacta'] = round(Descriptors.ExactMolWt(mol), 4)
    d['carga_formal'] = Chem.GetFormalCharge(mol)
    d['logp'] = round(Crippen.MolLogP(mol), 2)
    d['refractividad_molar'] = round(Crippen.MolMR(mol), 2)
    d['tpsa'] = round(Descriptors.TPSA(mol), 2)
    d['hbd'] = Lipinski.NumHDonors(mol)
    d['hba'] = Lipinski.NumHAcceptors(mol)
    d['enlaces_rotables'] = rdMolDescriptors.CalcNumRotatableBonds(mol)
    d['atomos_pesados'] = mol.GetNumHeavyAtoms()
    d['anillos'] = rdMolDescriptors.CalcNumRings(mol)
    d['anillos_aromaticos'] = rdMolDescriptors.CalcNumAromaticRings(mol)
    d['fraccion_csp3'] = round(rdMolDescriptors.CalcFractionCSP3(mol), 3)
    d['qed'] = round(QED.qed(mol), 3)
    d['estereocentros_cip'] = descriptores_cip(mol)
    d['inchi'] = Chem.MolToInchi(mol)
    d['inchikey'] = Chem.MolToInchiKey(mol)
    # Regla de los cinco de Lipinski (informativa: los neurotransmisores son
    # sustratos de transportadores, no farmacos orales)
    viol = sum([d['mw'] > 500, d['logp'] > 5, d['hbd'] > 5, d['hba'] > 10])
    d['lipinski_violaciones'] = viol
    d['cumple_lipinski'] = 'Si' if viol <= 1 else 'No'
    return mol, d


def dibujar(mol, ruta, leyenda='', ancho=500, alto=400):
    m = Chem.Mol(mol)
    AllChem.Compute2DCoords(m)
    drawer = rdMolDraw2D.MolDraw2DCairo(ancho, alto)
    opts = drawer.drawOptions()
    opts.addStereoAnnotation = True
    opts.bondLineWidth = 2
    rdMolDraw2D.PrepareAndDrawMolecule(drawer, m, legend=leyenda)
    drawer.FinishDrawing()
    with open(ruta, 'wb') as fh:
        fh.write(drawer.GetDrawingText())


def slug(texto):
    import re
    t = texto.lower()
    for a, b in zip('aaeiouunc', 'aaeiouunc'):
        pass
    t = (t.replace('á', 'a').replace('é', 'e').replace('í', 'i')
          .replace('ó', 'o').replace('ú', 'u').replace('ñ', 'n'))
    t = re.sub(r'[^a-z0-9]+', '_', t).strip('_')
    return t


def main():
    # 'notas' es la ultima columna y suele llevar comas. Si alguien edita el CSV
    # a mano sin entrecomillar ese campo, las comas sobrantes desbordan a columnas
    # extra y la nota se trunca en silencio. Aqui se detecta y se recompone.
    with open(MAESTRO, encoding='utf-8') as fh:
        lector = csv.DictReader(fh, restkey='__resto__')
        filas = []
        desbordadas = []
        for fila in lector:
            resto = fila.pop('__resto__', None)
            if resto:
                fila['notas'] = ','.join([fila.get('notas') or ''] + resto)
                desbordadas.append(fila['id'])
            filas.append(fila)
    if desbordadas:
        print('[aviso] campo notas sin entrecomillar en %d fila(s): %s'
              % (len(desbordadas), ', '.join(desbordadas)))
        print('        se han recompuesto para este calculo; conviene guardar el')
        print('        maestro desde Excel o con csv.writer para entrecomillarlas.')
    salida = []
    for r in filas:
        mol, d = analizar(r['smiles'])
        img = os.path.join(IMGDIR, '%s_%s.png' % (r['id'], slug(r['nombre'])))
        dibujar(mol, img, leyenda=r['nombre'])
        # forma predominante a pH 7.4
        carga_ph74 = ''
        if r.get('smiles_ph74'):
            m74 = Chem.MolFromSmiles(r['smiles_ph74'])
            if m74 is not None:
                carga_ph74 = Chem.GetFormalCharge(m74)
        fila = {
            'id': r['id'], 'nombre': r['nombre'], 'nombre_en': r['nombre_en'],
            'tema': r['tema'], 'familia': r['familia'], 'cas': r['cas'],
            'chembl_id': r.get('chembl_id', ''),
            'smiles_entrada': r['smiles'], 'smiles_ph74': r.get('smiles_ph74', ''),
            'carga_neta_ph74': carga_ph74,
        }
        fila.update(d)
        fila['imagen'] = os.path.relpath(img, BASE).replace('\\', '/')
        fila['notas'] = r['notas']
        salida.append(fila)
        print('[OK] %-14s %-28s %-12s MW %7.2f  logP %6.2f  TPSA %6.2f  %s'
              % (r['id'], r['nombre'][:28], d['formula'], d['mw'], d['logp'],
                 d['tpsa'], d['inchikey']))

    cols = list(salida[0].keys())
    ruta_csv = os.path.join(BASE, 'propiedades_qfdos.csv')
    with open(ruta_csv, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(salida)
    print('\n-> %s (%d estructuras)' % (ruta_csv, len(salida)))

    try:
        exportar_xlsx(salida)
    except ImportError:
        print('[aviso] openpyxl no disponible: se omite el XLSX')
    return salida


def exportar_xlsx(salida):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Propiedades'
    cols = [c for c in salida[0].keys() if c not in ('inchi',)]
    ws.append(cols)
    cab = PatternFill('solid', fgColor='1F3864')
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = cab
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in salida:
        ws.append([r[c] for c in cols])
    anchos = {'nombre': 28, 'nombre_en': 22, 'familia': 22, 'tema': 18,
              'smiles_entrada': 34, 'smiles_ph74': 34, 'smiles_canonico': 34,
              'inchikey': 30, 'notas': 60, 'estereocentros_cip': 18, 'imagen': 30}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(c, 14)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet('Estructuras')
    ws2.append(['id', 'nombre', 'formula', 'estructura'])
    for i, c in enumerate(['id', 'nombre', 'formula', 'estructura'], 1):
        ws2.cell(row=1, column=i).font = Font(bold=True, color='FFFFFF')
        ws2.cell(row=1, column=i).fill = cab
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 30
    ws2.column_dimensions['C'].width = 16
    ws2.column_dimensions['D'].width = 42
    fila = 2
    for r in salida:
        ws2.cell(row=fila, column=1, value=r['id'])
        ws2.cell(row=fila, column=2, value=r['nombre'])
        ws2.cell(row=fila, column=3, value=r['formula'])
        ruta = os.path.join(BASE, r['imagen'])
        if os.path.exists(ruta):
            img = XLImage(ruta)
            img.width, img.height = 250, 200
            ws2.row_dimensions[fila].height = 155
            ws2.add_image(img, 'D%d' % fila)
        fila += 1

    ruta = os.path.join(BASE, 'estructuras_qfdos.xlsx')
    wb.save(ruta)
    print('-> %s' % ruta)


if __name__ == '__main__':
    main()
